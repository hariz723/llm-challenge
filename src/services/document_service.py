import io
import uuid
from functools import lru_cache
from typing import List
import docx
import numpy as np
from pypdf import PdfReader
from huggingface_hub import InferenceClient
from fastapi import UploadFile, HTTPException
from ..schemas.auth import AuthenticatedUser
from ..core.logging import logger
from ..repository.document import DocumentRepository
from ..storage.azure_storage import AzureStorage
from ..schemas.document import DocumentSearchResponse, RAGChatResponse, RAGSource
from qdrant_client import QdrantClient, models
from sentence_transformers import SentenceTransformer
import core.constants as cons
from ..core.config import settings
import requests
from langfuse import observe, propagate_attributes
from ..core.langfuse_client import get_langfuse


@lru_cache(maxsize=1)
def get_hf_client() -> InferenceClient | None:
    token = settings.HUGGINGFACE_TOKEN
    if token:
        return InferenceClient(api_key=token)
    return None


embedding_model = SentenceTransformer(cons.EMBEDDING_MODEL_NAME)
EMBEDDING_DIMENSION = embedding_model.get_sentence_embedding_dimension()


def extract_text_from_file(content: bytes, filename: str) -> str:
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    try:
        if ext == "pdf":
            reader = PdfReader(io.BytesIO(content))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(p.strip() for p in pages if p.strip())

        if ext in ("docx", "doc"):
            doc = docx.Document(io.BytesIO(content))
            elements = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_text:
                        elements.append(" | ".join(row_text))
            return "\n".join(elements)

        # Fallback for plain text, markdown, code, json, csv, etc.
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="ignore")

    except Exception as e:
        logger.error(f"Error extracting text from {filename}: {e}")
        return ""


# Placeholder for text chunking (replace with actual implementation)
def chunk_text(text: str) -> List[str]:
    chunk_size = 700
    overlap = 120
    if not text.strip():
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end].strip())
        start += chunk_size - overlap
    return [chunk for chunk in chunks if chunk]


@observe(name="generate_embeddings", as_type="embedding")
def get_embeddings_batch(texts: List[str]) -> List[List[float]]:
    if not texts:
        return []

    lf_client = get_langfuse()
    client = get_hf_client()
    if client:
        try:
            all_vectors: List[List[float]] = []
            batch_size = 32
            for i in range(0, len(texts), batch_size):
                batch = texts[i : i + batch_size]
                emb = client.feature_extraction(
                    batch, model=settings.HF_EMBEDDING_MODEL
                )
                arr = emb if hasattr(emb, "shape") else np.array(emb)
                if arr.ndim == 3:
                    arr = np.mean(arr, axis=1)
                elif arr.ndim == 1:
                    arr = arr.reshape(1, -1)
                all_vectors.extend(arr.tolist())
            if lf_client:
                lf_client.update_current_span(
                    metadata={
                        "model": settings.HF_EMBEDDING_MODEL,
                        "count": len(texts),
                        "backend": "huggingface",
                    }
                )
            return [[float(x) for x in v] for v in all_vectors]
        except Exception as e:
            logger.warning(
                f"Hugging Face batch feature extraction failed ({e}). Falling back to local SentenceTransformer."
            )

    encoded = embedding_model.encode(texts, batch_size=32, show_progress_bar=False)
    if lf_client:
        lf_client.update_current_span(
            metadata={
                "model": cons.EMBEDDING_MODEL_NAME,
                "count": len(texts),
                "backend": "local_sentence_transformers",
            }
        )
    return [[float(x) for x in v] for v in encoded.tolist()]


def get_embedding(text: str) -> List[float]:
    results = get_embeddings_batch([text])
    return results[0] if results else []


class DocumentService:
    def __init__(
        self,
        document_repository: DocumentRepository,
        azure_storage: AzureStorage,
        qdrant_client: QdrantClient,
    ):
        self.document_repository = document_repository
        self.azure_storage = azure_storage
        self.qdrant_client = qdrant_client

    def _get_or_create_collection(self, collection_name: str):
        try:
            self.qdrant_client.get_collection(collection_name=collection_name)
            logger.info(f"Collection '{collection_name}' already exists.")
        except Exception:
            logger.info(f"Collection '{collection_name}' not found. Creating it.")
            self.qdrant_client.create_collection(
                collection_name=collection_name,
                vectors_config=models.VectorParams(
                    size=EMBEDDING_DIMENSION, distance=models.Distance.COSINE
                ),
            )

    def _get_trace_info(self) -> tuple[str | None, str | None]:
        trace_id = None
        trace_url = None
        client = get_langfuse()
        if client:
            try:
                trace_id = client.get_current_trace_id()
                if trace_id:
                    trace_url = client.get_trace_url(trace_id=trace_id)
            except Exception as exc:
                logger.debug(f"Could not retrieve Langfuse trace info: {exc}")
        return trace_id, trace_url

    @observe(name="retrieve_context", as_type="retriever")
    def _search_chunks(
        self, query: str, current_user: AuthenticatedUser, top_k: int = 5
    ) -> List[DocumentSearchResponse]:
        collection_name = str(current_user.id)
        lf_client = get_langfuse()
        try:
            self.qdrant_client.get_collection(collection_name=collection_name)
        except Exception:
            if lf_client:
                lf_client.update_current_span(
                    metadata={"collection_found": False, "hits_count": 0}
                )
            return []

        query_embedding = get_embedding(query)
        if hasattr(self.qdrant_client, "query_points"):
            response = self.qdrant_client.query_points(
                collection_name=collection_name,
                query=query_embedding,
                limit=top_k,
                with_payload=True,
            )
            hits = response.points
        else:
            hits = self.qdrant_client.search(
                collection_name=collection_name,
                query_vector=query_embedding,
                limit=top_k,
            )

        results = []
        for hit in hits:
            payload = hit.payload or {}
            results.append(
                DocumentSearchResponse(
                    document_id=payload["document_id"],
                    filename=payload["filename"],
                    blob_url=payload["blob_url"],
                    text=payload["text"],
                    score=hit.score,
                )
            )

        if lf_client:
            lf_client.update_current_span(
                metadata={
                    "collection": collection_name,
                    "top_k": top_k,
                    "hits_count": len(results),
                    "top_score": results[0].score if results else None,
                }
            )
        return results

    @observe(name="generate_answer_with_llm", as_type="generation")
    def _generate_answer_with_llm(
        self, query: str, search_results: List[DocumentSearchResponse]
    ) -> str | None:
        context_blocks = []
        for idx, result in enumerate(search_results, start=1):
            context_blocks.append(
                f"[Source {idx}] {result.filename}\n{result.text.strip()}"
            )

        prompt = (
            "Answer the user's question using only the provided sources. "
            "If the answer is not supported by the sources, say that clearly. "
            "Cite sources inline like [Source 1].\n\n"
            f"Question: {query}\n\n"
            f"Sources:\n\n{'\n\n'.join(context_blocks)}"
        )

        lf_client = get_langfuse()

        # 1. Try Hugging Face Inference if token is configured
        client = get_hf_client()
        if client:
            try:
                hf_model = settings.HF_CHAT_MODEL or "meta-llama/Llama-3.2-3B-Instruct"
                response = client.chat_completion(
                    messages=[
                        {
                            "role": "system",
                            "content": (
                                "You are an accurate and concise RAG assistant. "
                                "Answer the question strictly using the provided sources. "
                                "Cite sources like [Source 1]. If not found in sources, state that clearly."
                            ),
                        },
                        {"role": "user", "content": prompt},
                    ],
                    model=hf_model,
                    temperature=0.2,
                    max_tokens=512,
                )
                if response and response.choices:
                    content = response.choices[0].message.content
                    if content:
                        answer_text = content.strip()
                        if lf_client:
                            usage_details = None
                            if hasattr(response, "usage") and response.usage:
                                usage_details = {
                                    "input": getattr(response.usage, "prompt_tokens", 0)
                                    or 0,
                                    "output": getattr(
                                        response.usage, "completion_tokens", 0
                                    )
                                    or 0,
                                    "total": getattr(response.usage, "total_tokens", 0)
                                    or 0,
                                }
                            lf_client.update_current_generation(
                                model=hf_model,
                                model_parameters={
                                    "temperature": 0.2,
                                    "max_tokens": 512,
                                },
                                input=prompt,
                                output=answer_text,
                                usage_details=usage_details,
                            )
                        return answer_text
            except Exception as e:
                logger.error(f"Hugging Face chat completion failed: {e}")

        # 2. Try custom LLM_API_URL if configured
        if settings.LLM_API_URL and settings.LLM_MODEL:
            try:
                headers = {"Content-Type": "application/json"}
                if settings.LLM_API_KEY:
                    headers["Authorization"] = f"Bearer {settings.LLM_API_KEY}"

                response = requests.post(
                    settings.LLM_API_URL,
                    headers=headers,
                    json={
                        "model": settings.LLM_MODEL,
                        "messages": [
                            {
                                "role": "system",
                                "content": "You are a careful RAG assistant.",
                            },
                            {"role": "user", "content": prompt},
                        ],
                        "temperature": 0.2,
                    },
                    timeout=30,
                )
                response.raise_for_status()
                payload = response.json()
                raw_answer = payload["choices"][0]["message"]["content"].strip()
                if lf_client:
                    usage = payload.get("usage") or {}
                    usage_details = (
                        {
                            "input": usage.get("prompt_tokens", 0),
                            "output": usage.get("completion_tokens", 0),
                            "total": usage.get("total_tokens", 0),
                        }
                        if usage
                        else None
                    )
                    lf_client.update_current_generation(
                        model=settings.LLM_MODEL,
                        model_parameters={"temperature": 0.2},
                        input=prompt,
                        output=raw_answer,
                        usage_details=usage_details,
                    )
                return raw_answer
            except Exception as e:
                logger.error(f"Custom LLM API request failed: {e}")

        if lf_client:
            lf_client.update_current_generation(
                level="WARNING",
                status_message="LLM synthesis unavailable or returned empty; using fallback.",
            )

        return None

    @observe(name="generate_fallback_answer", as_type="generation")
    def _generate_fallback_answer(
        self, query: str, search_results: List[DocumentSearchResponse]
    ) -> str:
        lf_client = get_langfuse()
        if not search_results:
            msg = (
                "I couldn't find any relevant document context for that question yet. "
                "Upload a document first or try a more specific query."
            )
            if lf_client:
                lf_client.update_current_generation(
                    model="extractive_fallback",
                    input=query,
                    output=msg,
                    metadata={"results_count": 0},
                )
            return msg

        lines = []
        for idx, result in enumerate(search_results[:3], start=1):
            excerpt = " ".join(result.text.split())
            excerpt = excerpt[:280].rstrip()
            lines.append(f"[Source {idx}] {excerpt}")

        ans = (
            f"Based on the retrieved document chunks, the most relevant information for "
            f"'{query}' is:\n\n" + "\n\n".join(lines)
        )
        if lf_client:
            lf_client.update_current_generation(
                model="extractive_fallback",
                input=query,
                output=ans,
                metadata={
                    "strategy": "top_chunks_extraction",
                    "chunk_count": len(search_results[:3]),
                },
            )
        return ans

    @observe(name="document_upload_pipeline")
    async def upload_document(self, file: UploadFile, current_user: AuthenticatedUser):
        with propagate_attributes(
            user_id=str(current_user.id),
            tags=["upload", "ingestion"],
            metadata={"filename": file.filename},
        ):
            try:
                blob_url = await self.azure_storage.upload_file(file, str(uuid.uuid4()))

                await file.seek(0)
                content = await file.read()

                text = extract_text_from_file(content, file.filename)
                if not text.strip():
                    raise HTTPException(
                        status_code=400, detail="Could not extract text from file"
                    )

                chunks = chunk_text(text)
                collection_name = str(current_user.id)
                self._get_or_create_collection(collection_name)

                # Use the repository to create the document first to get its ID
                db_doc = await self.document_repository.create_document(
                    filename=file.filename,
                    user_id=current_user.id,
                    blob_url=blob_url,
                )

                embeddings = get_embeddings_batch(chunks)
                points = []
                for i, chunk in enumerate(chunks):
                    point_id = str(uuid.uuid4())
                    points.append(
                        models.PointStruct(
                            id=point_id,
                            vector=embeddings[i],
                            payload={
                                "document_id": db_doc.id,
                                "text": chunk,
                                "filename": file.filename,
                                "chunk_id": i,
                                "blob_url": blob_url,
                            },
                        )
                    )

                if points:
                    self.qdrant_client.upsert(
                        collection_name=collection_name,
                        points=points,
                        wait=True,
                    )

                return db_doc

            except HTTPException:
                raise
            except Exception as e:
                logger.error(f"Error uploading document: {e}")
                raise HTTPException(status_code=500, detail="Error processing document")

    @observe(name="search_documents_pipeline")
    async def search_documents(
        self, query: str, current_user: AuthenticatedUser
    ) -> List[DocumentSearchResponse]:
        with propagate_attributes(
            user_id=str(current_user.id),
            tags=["vector_search"],
            metadata={"query": query},
        ):
            try:
                return self._search_chunks(
                    query=query, current_user=current_user, top_k=5
                )
            except Exception as e:
                logger.error(f"Error searching documents: {e}")
                raise HTTPException(status_code=500, detail="Error searching documents")

    @observe(name="rag_chat_pipeline")
    async def answer_question(
        self, query: str, current_user: AuthenticatedUser, top_k: int = 5
    ) -> RAGChatResponse:
        with propagate_attributes(
            user_id=str(current_user.id),
            tags=["rag", "chat"],
            metadata={"top_k": top_k, "query": query},
        ):
            try:
                search_results = self._search_chunks(
                    query=query, current_user=current_user, top_k=top_k
                )
                answer = self._generate_answer_with_llm(query, search_results)
                if not answer:
                    answer = self._generate_fallback_answer(query, search_results)

                trace_id, trace_url = self._get_trace_info()

                return RAGChatResponse(
                    answer=answer,
                    sources=[
                        RAGSource(
                            document_id=result.document_id,
                            filename=result.filename,
                            blob_url=result.blob_url,
                            text=result.text,
                            score=result.score,
                        )
                        for result in search_results
                    ],
                    trace_id=trace_id,
                    trace_url=trace_url,
                )
            except requests.RequestException as exc:
                logger.warning(f"LLM generation failed, using fallback answer: {exc}")
                search_results = self._search_chunks(
                    query=query, current_user=current_user, top_k=top_k
                )
                fallback_ans = self._generate_fallback_answer(query, search_results)
                trace_id, trace_url = self._get_trace_info()
                return RAGChatResponse(
                    answer=fallback_ans,
                    sources=[
                        RAGSource(
                            document_id=result.document_id,
                            filename=result.filename,
                            blob_url=result.blob_url,
                            text=result.text,
                            score=result.score,
                        )
                        for result in search_results
                    ],
                    trace_id=trace_id,
                    trace_url=trace_url,
                )
            except Exception as e:
                logger.error(f"Error answering question: {e}")
                raise HTTPException(status_code=500, detail="Error generating answer")
